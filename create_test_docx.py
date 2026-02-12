#!/usr/bin/env python3
"""Create a test DOCX file with text on specific pages."""

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_BREAK

doc = Document()

# Page 1
doc.add_heading("Test Document - Page 1", 0)
doc.add_paragraph("This is the first page. It contains the word 'TARGET' here.")
doc.add_paragraph("Some more text to fill the page.")

# Force page break to page 2
doc.add_page_break()

# Page 2
doc.add_heading("Page 2", 1)
doc.add_paragraph("This is the second page. No target word here on this page.")
doc.add_paragraph("Just some regular text.")

# Force page break to page 3
doc.add_page_break()

# Page 3
doc.add_heading("Page 3", 1)
doc.add_paragraph("Third page with 'TARGET' appearing again here.")
doc.add_paragraph("And another occurrence of TARGET on this page.")
doc.add_paragraph("More content to test.")

# Force page break to page 4
doc.add_page_break()

# Page 4
doc.add_heading("Page 4", 1)
doc.add_paragraph("Final page. No target word here.")

doc.save("test_document.docx")
print("Created test_document.docx with 'TARGET' on pages 1 and 3")
